"""
build/to_pptx.py
================
Compile source/*.hymn into one PowerPoint deck per hymn:

  dist/pptx/NNN.pptx   -- title slide plus one slide per stanza

Every deck is built on template/, packed by build.pptx_template. That template
carries one slide master ("Reformation Hymnal") and two layouts, "Hymn Title"
and "Hymn Stanza". Slides store text and nothing else, so restyling the whole
library is a template edit and a rebuild rather than 3,719 slide edits.

Why the slides carry a scale instead of a point size
----------------------------------------------------
The first version wrote the chosen point size onto every paragraph. It looked
right, but a hard ``sz`` on the slide beats anything the destination layout
says, so pasting a stanza into another deck kept our size until someone
right-clicked and chose Reset Slide.

So the slides set no size at all. The size comes from the layout placeholder,
and the slide carries only ``<a:normAutofit fontScale="...">`` -- the same
shrink-on-overflow PowerPoint applies from the Format Shape pane. Text then
inherits the destination's size on paste and our scale rides along on top of
it, proportionally.

The scale has to be written, not left for PowerPoint to work out. PowerPoint
caches ``fontScale`` in the file and only recomputes it when the text reflows,
so a deck shipped with a bare ``<a:normAutofit/>`` renders at full size and
overflows until a human edits the box. Computing it is the same fitting work
we were already doing: the layout default is the largest size in the ladder,
and the scale is whatever fraction of it the stanza actually needs.

Fitting itself measures real text widths in the real font (Pillow, which
arrives with python-pptx) and word-wraps the way PowerPoint does. Counting
source lines is not enough: hymn 260's chorus is four lines but its longest
line is 85 characters, and at the size four lines suggests it wraps to seven
and runs off the bottom of the slide.

One size is used for the whole hymn, taken from whichever section needs the
smallest. A song that changes text size between verse and chorus reads as
broken from the pews. Across the 700 hymns this lands 683 at 46pt or larger
and bottoms out at 30pt for hymn 255, whose verses run twelve lines.

Usage:
    python -m build.to_pptx              # all 700
    python -m build.to_pptx 1 255 260    # just these hymn numbers
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from build.model import Hymn, Section
from build.pptx_template import pack

ROOT = Path(__file__).parent.parent
SOURCE = ROOT / "source"
OUT = ROOT / "dist" / "pptx"
BUNDLE = ROOT / "dist" / "ReformationHymnal-pptx.zip"

# The theme asks for Cambria and Calibri. Neither is redistributable, so a CI
# runner measures with Google's metric-compatible pair instead. Advance widths
# match closely enough that the fitted size is identical for all 700 hymns,
# verified by running the fitter over the corpus with both sets.
FONT_MAJOR = ("Cambria", "Caladea")
FONT_MINOR = ("Calibri", "Carlito")

LAYOUT_TITLE = "Hymn Title"
LAYOUT_STANZA = "Hymn Stanza"

# Placeholder idx values, fixed by template/. See build/make_template.py.
PH_TITLE, PH_NUMBER, PH_ATTRIBUTION = 0, 1, 2
PH_STANZA, PH_REFERENCE, PH_END = 1, 2, 3

END_GLYPH = "* * *"                      # parity with SoftProjector and FreeShow

# Sizes are tried from the layout default downward in these steps, so a change
# to the template default retunes the whole library without touching this file.
SIZE_STEP = 4
SIZE_FLOOR = 22

# PowerPoint's default single line spacing, near enough for fitting.
LINE_SPACING = 1.2

# The template's default text insets, 0.05in a side.
INSET_PT = 7.2


# ---------------------------------------------------------------------------
# Text measurement
# ---------------------------------------------------------------------------

#: Size the fonts are loaded at. Advance widths scale linearly, so one
#: measurement divided by this gives width per point of font size.
_MEASURE_PX = 200


def _load_font(*names: str):
    """Resolve a theme font to something Pillow can measure.

    Names are tried in order, so a host can supply a metric-compatible stand-in
    for a font it is not licensed to ship. fc-match always returns *something*,
    so the family it returns is checked against the name asked for: falling
    through to an arbitrary face would change every fitted size silently, and
    that is worth a loud warning instead.
    """
    import subprocess

    from PIL import ImageFont

    for name in names:
        try:
            matched = subprocess.run(
                ["fc-match", "-f", "%{file}\t%{family}", name],
                capture_output=True, text=True, check=True,
            ).stdout.split("\t")
        except (OSError, subprocess.CalledProcessError):
            break
        if len(matched) == 2 and name.lower() in matched[1].lower():
            return ImageFont.truetype(matched[0], _MEASURE_PX)
    print(f"warning: none of {', '.join(names)} is installed; fitted sizes "
          f"will not match a machine that has them", file=sys.stderr)
    from PIL import ImageFont as _IF
    return _IF.load_default(_MEASURE_PX)


def _ladder(base_pt: int) -> list[int]:
    return list(range(base_pt, SIZE_FLOOR - 1, -SIZE_STEP)) or [base_pt]


class Fitter:
    """Picks the largest size at or below the layout default that fits."""

    def __init__(self, font, layout, idx: int):
        ph = _placeholder(layout.placeholders, idx)
        self.font = font
        self.width = Emu(ph.width).pt - INSET_PT
        self.height = Emu(ph.height).pt - INSET_PT
        self.base = _default_pt(ph, layout.name, idx)
        self.sizes = _ladder(self.base)

    def _width_per_pt(self, text: str) -> float:
        return self.font.getlength(text) / _MEASURE_PX

    def _visual_lines(self, line: str, size: int) -> int:
        """Greedy word wrap, which is how PowerPoint breaks a paragraph."""
        words = line.split()
        if not words:
            return 1
        count, current = 1, ""
        for word in words:
            trial = f"{current} {word}".strip()
            if self._width_per_pt(trial) * size <= self.width:
                current = trial
            else:
                count, current = count + 1, word
        return count

    def fit(self, lines: list[str]) -> int:
        for size in self.sizes:
            total = sum(self._visual_lines(line, size) for line in lines)
            if total * size * LINE_SPACING <= self.height:
                return size
        return self.sizes[-1]


# ---------------------------------------------------------------------------
# Slide plumbing
# ---------------------------------------------------------------------------

def _placeholder(shapes, idx):
    """Look up by placeholder idx.

    python-pptx keys a slide's ``placeholders`` by idx but a layout's by
    position, which is a quiet source of wrong-box bugs. Match explicitly.
    """
    return next(s for s in shapes if s.placeholder_format.idx == idx)


def _default_pt(placeholder, layout_name: str, idx: int) -> int:
    """The size a layout placeholder hands down, in points."""
    path = f"./{qn('p:txBody')}/{qn('a:lstStyle')}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}"
    defrpr = placeholder._element.find(path)
    size = None if defrpr is None else defrpr.get("sz")
    if size is None:
        raise SystemExit(
            f"template layout {layout_name!r} placeholder idx {idx} has no "
            f"explicit size in its lstStyle; the slide scale is a fraction of "
            f"that size, so it cannot be computed without one"
        )
    return int(size) // 100


def _fill(shapes, idx: int, lines: list[str]) -> None:
    """Put text in a placeholder and nothing else.

    Deliberately sets no run or paragraph size: anything written here would
    outrank the destination layout when the slide is pasted elsewhere.
    """
    frame = _placeholder(shapes, idx).text_frame
    frame.text = lines[0]
    for line in lines[1:]:
        frame.add_paragraph().text = line


def _scale(shapes, idx: int, fitted_pt: int, base_pt: int) -> None:
    """Shrink a placeholder's inherited size by writing PowerPoint's own
    autofit scale. Omitted at 100%, which is what PowerPoint does too."""
    body_pr = _placeholder(shapes, idx).text_frame._bodyPr
    for tag in ("a:normAutofit", "a:spAutoFit", "a:noAutofit"):
        existing = body_pr.find(qn(tag))
        if existing is not None:
            body_pr.remove(existing)
    autofit = body_pr.makeelement(qn("a:normAutofit"), {})
    if fitted_pt < base_pt:
        autofit.set("fontScale", str(round(fitted_pt / base_pt * 100000)))
    body_pr.append(autofit)


def _drop(shapes, idx: int) -> None:
    """Remove an unused placeholder rather than leaving it empty.

    An empty placeholder shows nothing in presentation view, but it does show
    prompt text while editing, and anyone pulling a slide into their own deck
    will be editing.
    """
    shape = _placeholder(shapes, idx)
    shape._element.getparent().remove(shape._element)


def _label(section: Section, verse_num: int, verse_count: int) -> str:
    if section.kind == "verse":
        return f"Verse {verse_num} of {verse_count}"
    return section.kind.capitalize()


# ---------------------------------------------------------------------------

def build_deck(hymn: Hymn, template: Path, fonts: dict) -> Presentation:
    prs = Presentation(str(template))
    layouts = {l.name: l for l in prs.slide_masters[0].slide_layouts}
    title_layout, stanza_layout = layouts[LAYOUT_TITLE], layouts[LAYOUT_STANZA]

    title_fit = Fitter(fonts["major"], title_layout, PH_TITLE)
    stanza_fit = Fitter(fonts["minor"], stanza_layout, PH_STANZA)

    slide = prs.slides.add_slide(title_layout).shapes
    _fill(slide, PH_NUMBER, [f"Hymn {hymn.number}"])
    _fill(slide, PH_TITLE, [hymn.title])
    _scale(slide, PH_TITLE, title_fit.fit([hymn.title]), title_fit.base)
    attribution = " / ".join(x for x in (hymn.author, hymn.composer) if x)
    if attribution:
        _fill(slide, PH_ATTRIBUTION, [attribution])
    else:
        _drop(slide, PH_ATTRIBUTION)

    size = min(stanza_fit.fit(s.lines) for s in hymn.sections)
    verse_count = sum(1 for s in hymn.sections if s.kind == "verse")

    verse_num = 0
    for i, section in enumerate(hymn.sections):
        if section.kind == "verse":
            verse_num += 1
        shapes = prs.slides.add_slide(stanza_layout).shapes
        _fill(shapes, PH_STANZA, section.lines)
        _scale(shapes, PH_STANZA, size, stanza_fit.base)
        _fill(shapes, PH_REFERENCE,
              [f"{hymn.number} · {_label(section, verse_num, verse_count)}"])
        if i == len(hymn.sections) - 1:
            _fill(shapes, PH_END, [END_GLYPH])
        else:
            _drop(shapes, PH_END)

    return prs


def main(argv: list[str]) -> int:
    template = pack()                     # always build from current template/
    fonts = {"major": _load_font(*FONT_MAJOR), "minor": _load_font(*FONT_MINOR)}

    if argv:
        wanted = {int(a) for a in argv}
        paths = [p for p in sorted(SOURCE.glob("*.hymn")) if int(p.stem) in wanted]
        missing = wanted - {int(p.stem) for p in paths}
        if missing:
            print(f"no such hymn: {sorted(missing)}", file=sys.stderr)
            return 1
    else:
        paths = sorted(SOURCE.glob("*.hymn"))

    OUT.mkdir(parents=True, exist_ok=True)
    slides = 0
    for path in paths:
        hymn = Hymn.load(path)
        build_deck(hymn, template, fonts).save(str(OUT / f"{hymn.number:03d}.pptx"))
        slides += len(hymn.sections) + 1

    print(f"Written {OUT}/  ({len(paths)} decks, {slides} slides)")

    # One zip is the release asset; 700 loose files is not.
    with zipfile.ZipFile(BUNDLE, "w", zipfile.ZIP_DEFLATED) as bundle:
        for deck in sorted(OUT.glob("*.pptx")):
            info = zipfile.ZipInfo(deck.name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            bundle.writestr(info, deck.read_bytes())
    print(f"Written {BUNDLE}  ({BUNDLE.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
